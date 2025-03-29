import os

import pandas as pd
from openpyxl.reader.excel import load_workbook

from .excel_writer_charts import ExcelWriterCharts


class ExcelWriter:
    def __init__(self, model_obj, solver):
        self.data_path = os.getcwd()
        self.results_file_path = self.data_path + "\output.xlsx"
        self.model_obj = model_obj
        self.solver = solver
        self.add_charts = ExcelWriterCharts(self.results_file_path)

    def write_results(self):
        results_data = []
        for c in self.model_obj.all_courses:
            for s in self.model_obj.all_students:
                if self.solver.value(self.model_obj.allocation[s][c]):
                    results_data.append([
                        self.model_obj.courses[c].course_name,
                        self.model_obj.courses[c].course_id,
                        self.model_obj.students[s].fullname,
                        self.model_obj.students[s].student_id,
                        self.model_obj.students[s].semester,
                        self.model_obj.students[s].gpa,
                        self.model_obj.students[s].preferences[self.model_obj.courses[c].course_id],
                        self.model_obj.students[s].is_obligated,
                        self.model_obj.students[s].passed_courses_on_sem8,
                        self.model_obj.students[s].courses_needed_remaining,
                        self.model_obj.students[s].choices_remaining,
                        sum(self.solver.value(self.model_obj.allocation[s][c]) for c in self.model_obj.all_courses)
                    ])

        df = pd.DataFrame(results_data, columns = [
            "Course Name",
            "Course ID",
            "Fullname",
            "AEM",
            "Semester",
            "GPA",
            "Preference",
            "Obligated",
            "Passed Courses On Sem8",
            "Courses Needed Remaining",
            "Choices Remaining",
            "Assigned Courses on Student"
        ])

        try:
            print(f"Assignment results are saved to {self.results_file_path} in sheet: Assignments.")
            df.to_excel(self.results_file_path, sheet_name="Assignments", index = False)
            self.write_results_per_course()
            self.add_charts.add_courses_sat_bar_chart()
        except Exception as e:
            print("An error occurred while writing to the assignment results excel file. \n", e)

        return df


    def write_sat_preferences(self):
        """This method collects data for each student to compare the satisfaction of their top 6 preferences
        with the courses they were assigned."""

        preferences_met_data = []
        top_6_preferences_satisfaction_ratios = []
        top_preferences = 6

        for s in self.model_obj.all_students:
            assigned_courses = [
                self.model_obj.courses[c].course_id for c in self.model_obj.all_courses
                if self.solver.value(self.model_obj.allocation[s][c])
            ]

            all_assigned_courses = len(assigned_courses)

            assigned_preferences = [
                self.model_obj.students[s].preferences[self.model_obj.courses[c].course_id]
                for c in self.model_obj.all_courses
                if self.solver.value(self.model_obj.allocation[s][c])
            ]

            all_assigned_preferences = len(assigned_preferences)

            # calculating how many of the courses that the student is assigned, they are in their top 6
            top_6_preferences_assigned = 0

            if all_assigned_courses > 0:
                for course_id in assigned_courses:
                    if (course_id in self.model_obj.students[s].preferences and
                            self.model_obj.students[s].preferences[course_id] <= top_preferences):
                        top_6_preferences_assigned += 1


            assigned_courses_string = ", ".join(map(str, assigned_courses))
            assigned_preferences_string = ", ".join(map(str, assigned_preferences))
            top_6_preferences_satisfaction_ratio = top_6_preferences_assigned / all_assigned_courses * 100

            top_6_preferences_satisfaction_ratios.append(top_6_preferences_satisfaction_ratio)

            preferences_met_data.append([
                self.model_obj.students[s].fullname,
                self.model_obj.students[s].student_id,
                self.model_obj.students[s].semester,
                self.model_obj.students[s].gpa,
                assigned_courses_string,
                assigned_preferences_string,
                all_assigned_courses,
                top_6_preferences_satisfaction_ratio
            ])

        df = pd.DataFrame(preferences_met_data, columns = [
            "Fullname",
            "AEM",
            "Semester",
            "GPA",
            "IDs of Assigned Courses on Student",
            "Assigned Preferences",
            "# of Assigned Courses on Student",
            "Top 6 Preferences Satisfaction Ratio on Student (%)"
        ])

        try:
            df_no_duplicates = df.drop_duplicates()

            with pd.ExcelWriter(
                    self.results_file_path,
                    engine="openpyxl",
                    mode="a" if os.path.exists(self.results_file_path) else "w",
                    if_sheet_exists="replace"
            ) as writer:
                df_no_duplicates.to_excel(writer, index=False, sheet_name="Preferences Satisfaction")

            wb = load_workbook(self.results_file_path)
            wb.save(self.results_file_path)
            self.add_charts.add_top_6_preferences_sat_pie_chart(top_6_preferences_satisfaction_ratios)
            print(
                f"Top 6 preferences satisfaction results are saved to {self.results_file_path} "
                f"in sheet: Preferences Satisfaction.")
        except Exception as e:
            print("An error occurred while writing to the top 6 preferences satisfaction results excel file. \n", e)


    def write_results_per_course(self):
        """This method creates a new sheet in the excel file and lists the students full names assigned to each course."""
        courses_data = {self.model_obj.courses[c].course_name: [] for c in self.model_obj.all_courses}

        for c in self.model_obj.all_courses:
            for s in self.model_obj.all_students:
                if self.solver.value(self.model_obj.allocation[s][c]):
                    courses_data[self.model_obj.courses[c].course_name].append(self.model_obj.students[s].fullname)

        max_length = max(len(students) for students in courses_data.values())

        # dictionary values need to have the same length
        for course_name in courses_data:
            while len(courses_data[course_name]) < max_length:
                courses_data[course_name].append(None)

        df = pd.DataFrame(courses_data)

        try:
            with pd.ExcelWriter(
                    self.results_file_path,
                    engine="openpyxl",
                    mode="a",
                    if_sheet_exists="replace"
            ) as writer:
                df.to_excel(writer, index=False, sheet_name="Course Results")

            print(f"Courses results are saved to {self.results_file_path} in sheet: Course Results.")
        except Exception as e:
            print("An error occurred while writing to the courses results excel file. \n", e)